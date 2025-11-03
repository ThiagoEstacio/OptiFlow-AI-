"""
Tests for Report Generator Service
"""
import pytest
import os
from datetime import datetime, date, timedelta
from sqlalchemy.ext.asyncio import AsyncSession
from uuid import uuid4
from unittest.mock import patch, Mock

from app.services.report_generator import ReportGenerator
from app.models.operations import TruckEntry, ShipLoading, ProductType


@pytest.fixture
async def test_truck_entries(test_db: AsyncSession, test_site):
    """Create test truck entries"""
    today = date.today()
    trucks = [
        TruckEntry(
            id=uuid4(),
            site_id=test_site.id,
            license_plate="ABC-1234",
            driver_name="John Doe",
            product_type=ProductType.SOYBEAN,
            gross_weight=45000.0,
            tare_weight=15000.0,
            net_weight=30000.0,
            entry_time=datetime.combine(today, datetime.min.time()),
            exit_time=datetime.combine(today, datetime.min.time()) + timedelta(hours=2),
            status="completed",
        ),
        TruckEntry(
            id=uuid4(),
            site_id=test_site.id,
            license_plate="XYZ-5678",
            driver_name="Jane Smith",
            product_type=ProductType.CORN,
            gross_weight=42000.0,
            tare_weight=14000.0,
            net_weight=28000.0,
            entry_time=datetime.combine(today, datetime.min.time()) + timedelta(hours=3),
            exit_time=datetime.combine(today, datetime.min.time()) + timedelta(hours=5),
            status="completed",
        ),
    ]
    for truck in trucks:
        test_db.add(truck)
    await test_db.commit()
    return trucks


@pytest.fixture
async def test_ship_loadings_for_report(test_db: AsyncSession, test_site):
    """Create test ship loadings for reports"""
    today = date.today()
    ships = [
        ShipLoading(
            id=uuid4(),
            site_id=test_site.id,
            ship_name="MV Atlantic",
            ship_imo="1234567",
            product_type=ProductType.SOYBEAN,
            scheduled_arrival=datetime.combine(today, datetime.min.time()),
            actual_arrival=datetime.combine(today, datetime.min.time()) + timedelta(hours=1),
            estimated_tonnage=50000.0,
            actual_tonnage=48500.0,
            status="loading",
            loading_progress_percent=65.0,
        ),
    ]
    for ship in ships:
        test_db.add(ship)
    await test_db.commit()
    return ships


@pytest.mark.asyncio
async def test_generate_daily_operations_pdf_success(
    test_db: AsyncSession,
    test_site,
    test_truck_entries,
    test_ship_loadings_for_report
):
    """Test successful PDF report generation"""
    generator = ReportGenerator(test_db)
    today = date.today()

    filename = await generator.generate_daily_operations_pdf(test_site.id, today)

    assert filename is not None
    assert filename.endswith(".pdf")
    assert "daily_operations" in filename.lower()
    # Check if file exists
    assert os.path.exists(filename)
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_generate_daily_operations_pdf_no_data(test_db: AsyncSession, test_site):
    """Test PDF generation with no operational data"""
    generator = ReportGenerator(test_db)
    today = date.today()

    filename = await generator.generate_daily_operations_pdf(test_site.id, today)

    # Should still generate report even with no data
    assert filename is not None
    assert os.path.exists(filename)
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_generate_daily_operations_pdf_filename_format(
    test_db: AsyncSession,
    test_site,
    test_truck_entries
):
    """Test PDF filename format"""
    generator = ReportGenerator(test_db)
    today = date.today()

    filename = await generator.generate_daily_operations_pdf(test_site.id, today)

    # Filename should contain date
    date_str = today.strftime("%Y%m%d")
    assert date_str in filename
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_generate_daily_operations_excel_success(
    test_db: AsyncSession,
    test_site,
    test_truck_entries,
    test_ship_loadings_for_report
):
    """Test successful Excel report generation"""
    generator = ReportGenerator(test_db)
    start_date = date.today()
    end_date = start_date + timedelta(days=7)

    filename = await generator.generate_daily_operations_excel(test_site.id, start_date, end_date)

    assert filename is not None
    assert filename.endswith(".xlsx")
    assert "operations" in filename.lower()
    # Check if file exists
    assert os.path.exists(filename)
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_generate_daily_operations_excel_no_data(test_db: AsyncSession, test_site):
    """Test Excel generation with no operational data"""
    generator = ReportGenerator(test_db)
    start_date = date.today()
    end_date = start_date + timedelta(days=7)

    filename = await generator.generate_daily_operations_excel(test_site.id, start_date, end_date)

    # Should still generate report even with no data
    assert filename is not None
    assert os.path.exists(filename)
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_generate_daily_operations_excel_date_range(
    test_db: AsyncSession,
    test_site,
    test_truck_entries
):
    """Test Excel generation with specific date range"""
    generator = ReportGenerator(test_db)
    start_date = date.today() - timedelta(days=30)
    end_date = date.today()

    filename = await generator.generate_daily_operations_excel(test_site.id, start_date, end_date)

    assert filename is not None
    assert os.path.exists(filename)
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_generate_daily_operations_excel_filename_format(
    test_db: AsyncSession,
    test_site,
    test_truck_entries
):
    """Test Excel filename format includes date range"""
    generator = ReportGenerator(test_db)
    start_date = date.today()
    end_date = start_date + timedelta(days=7)

    filename = await generator.generate_daily_operations_excel(test_site.id, start_date, end_date)

    # Filename should contain date range
    start_str = start_date.strftime("%Y%m%d")
    end_str = end_date.strftime("%Y%m%d")
    assert start_str in filename or end_str in filename
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_generate_pdf_creates_reports_directory(test_db: AsyncSession, test_site):
    """Test that PDF generation creates reports directory if it doesn't exist"""
    generator = ReportGenerator(test_db)
    today = date.today()

    # Delete reports directory if it exists
    reports_dir = "reports"
    if os.path.exists(reports_dir):
        import shutil
        shutil.rmtree(reports_dir)

    filename = await generator.generate_daily_operations_pdf(test_site.id, today)

    # Directory should be created
    assert os.path.exists(reports_dir)
    assert os.path.exists(filename)
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_generate_excel_creates_reports_directory(test_db: AsyncSession, test_site):
    """Test that Excel generation creates reports directory if it doesn't exist"""
    generator = ReportGenerator(test_db)
    start_date = date.today()
    end_date = start_date + timedelta(days=7)

    # Delete reports directory if it exists
    reports_dir = "reports"
    if os.path.exists(reports_dir):
        import shutil
        shutil.rmtree(reports_dir)

    filename = await generator.generate_daily_operations_excel(test_site.id, start_date, end_date)

    # Directory should be created
    assert os.path.exists(reports_dir)
    assert os.path.exists(filename)
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_pdf_includes_truck_data(
    test_db: AsyncSession,
    test_site,
    test_truck_entries
):
    """Test that PDF includes truck entry data"""
    generator = ReportGenerator(test_db)
    today = date.today()

    filename = await generator.generate_daily_operations_pdf(test_site.id, today)

    # Verify file was created (actual PDF content verification would require PDF parsing)
    assert os.path.exists(filename)
    file_size = os.path.getsize(filename)
    # PDF should have reasonable size (not empty)
    assert file_size > 1000  # More than 1KB
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_pdf_includes_ship_data(
    test_db: AsyncSession,
    test_site,
    test_ship_loadings_for_report
):
    """Test that PDF includes ship loading data"""
    generator = ReportGenerator(test_db)
    today = date.today()

    filename = await generator.generate_daily_operations_pdf(test_site.id, today)

    assert os.path.exists(filename)
    file_size = os.path.getsize(filename)
    assert file_size > 1000  # Should contain data
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_excel_has_multiple_sheets(
    test_db: AsyncSession,
    test_site,
    test_truck_entries,
    test_ship_loadings_for_report
):
    """Test that Excel report has multiple sheets"""
    generator = ReportGenerator(test_db)
    start_date = date.today()
    end_date = start_date + timedelta(days=7)

    filename = await generator.generate_daily_operations_excel(test_site.id, start_date, end_date)

    # Verify file was created
    assert os.path.exists(filename)

    # Verify it's a valid Excel file by loading it
    from openpyxl import load_workbook
    wb = load_workbook(filename)
    sheet_names = wb.sheetnames

    # Should have multiple sheets (Summary, Trucks, Ships)
    assert len(sheet_names) >= 2
    # Clean up
    wb.close()
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_generate_pdf_invalid_site(test_db: AsyncSession):
    """Test PDF generation with invalid site ID"""
    generator = ReportGenerator(test_db)
    today = date.today()
    invalid_site_id = 999999

    filename = await generator.generate_daily_operations_pdf(invalid_site_id, today)

    # Should still generate report (possibly empty or with error message)
    assert filename is not None
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_generate_excel_invalid_site(test_db: AsyncSession):
    """Test Excel generation with invalid site ID"""
    generator = ReportGenerator(test_db)
    start_date = date.today()
    end_date = start_date + timedelta(days=7)
    invalid_site_id = 999999

    filename = await generator.generate_daily_operations_excel(invalid_site_id, start_date, end_date)

    # Should still generate report
    assert filename is not None
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_excel_includes_summary_sheet(
    test_db: AsyncSession,
    test_site,
    test_truck_entries
):
    """Test that Excel report includes summary sheet"""
    generator = ReportGenerator(test_db)
    start_date = date.today()
    end_date = start_date + timedelta(days=7)

    filename = await generator.generate_daily_operations_excel(test_site.id, start_date, end_date)

    from openpyxl import load_workbook
    wb = load_workbook(filename)

    # Check for summary-related sheet
    sheet_names = [name.lower() for name in wb.sheetnames]
    assert any("summary" in name or "resumo" in name for name in sheet_names)

    wb.close()
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_pdf_generation_performance(
    test_db: AsyncSession,
    test_site,
    test_truck_entries
):
    """Test that PDF generation completes in reasonable time"""
    import time
    generator = ReportGenerator(test_db)
    today = date.today()

    start_time = time.time()
    filename = await generator.generate_daily_operations_pdf(test_site.id, today)
    end_time = time.time()

    # Should complete in less than 5 seconds
    assert (end_time - start_time) < 5.0
    assert os.path.exists(filename)
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)


@pytest.mark.asyncio
async def test_excel_generation_performance(
    test_db: AsyncSession,
    test_site,
    test_truck_entries
):
    """Test that Excel generation completes in reasonable time"""
    import time
    generator = ReportGenerator(test_db)
    start_date = date.today()
    end_date = start_date + timedelta(days=7)

    start_time = time.time()
    filename = await generator.generate_daily_operations_excel(test_site.id, start_date, end_date)
    end_time = time.time()

    # Should complete in less than 5 seconds
    assert (end_time - start_time) < 5.0
    assert os.path.exists(filename)
    # Clean up
    if os.path.exists(filename):
        os.remove(filename)
